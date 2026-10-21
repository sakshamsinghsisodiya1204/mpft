"""Adds a slide on role-specific views, scenarios and the graph axis to the submission presentation."""
import copy
from pathlib import Path
import pptx
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt
from PIL import Image

F, A = "docs/submission/MPFT_Presentation.pptx", Path("docs/submission/assets")
pr = pptx.Presentation(F)
if any("Views and Scenarios" in sh.text_frame.text for s in pr.slides for sh in s.shapes if sh.has_text_frame):
    raise SystemExit("already updated")
ref = pr.slides[5]
new = pr.slides.add_slide(ref.slide_layout)
for sh in ref.shapes:
    if sh.name in ("Rectangle 7", "Rectangle 1", "Rectangle 2"):
        new.shapes._spTree.append(copy.deepcopy(sh._element))
for sh in new.shapes:
    if sh.name == "Rectangle 7":
        sh.text_frame.paragraphs[0].runs[0].text = "Views and Scenarios"
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r._r.getparent().remove(r._r)
    if sh.name == "Rectangle 2":
        sh.text_frame.paragraphs[0].runs[0].text = "35 seconds"


def pic(png, x, w, y=1250000):
    h = int(w * Image.open(png).size[1] / Image.open(png).size[0])
    new.shapes.add_picture(str(png), Emu(x), Emu(y), Emu(w), Emu(h))
    return y + h


def label(text, x, y, w, size=11, bold=False, color=RGBColor(0x16, 0x32, 0x4F)):
    tb = new.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(300000))
    tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold, r.font.color.rgb = text, Pt(size), bold, color


W = 4000000
b1 = pic(A / "scn_sender.png", 523875, W)
b2 = pic(A / "scn_receiver.png", 523875 + W + 150000, W)
y = max(b1, b2) + 40000
label("Sender: Delivered, send another", 523875, y, W, bold=True)
label("Receiver: verified, Download", 523875 + W + 150000, y, W, bold=True)
label("Eight scenarios: clean, slow, asymmetric, lossy, heavy loss, path failure, path flap, total failure (nothing saved, no download).", 523875, y + 330000, 8150000)
label("Graphs now use real time: throughput from chunk completions, millisecond axis under 2 s.", 523875, y + 800000, 8150000)
new.notes_slide.notes_text_frame.text = "Same session, two roles: sender sees delivery, receiver sees verification and Download. Pick a scenario in the Send form to show success and failure (35 s)."
lst = pr.slides._sldIdLst
el = lst[-1]
lst.remove(el)
lst.insert(6, el)
pr.save(F)
print("saved", len(pr.slides))
