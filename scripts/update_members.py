"""Adds the third team member to the submission report and slides."""
import copy
import docx, pptx
from pptx.util import Emu

R = "docs/submission/MPFT_Mini_Project_Report.docx"
NEW = ("Anupria Gupta", "1262241349")
d = docx.Document(R)


def add_row(t, values):
    row = copy.deepcopy(t.rows[-1]._tr)
    t._tbl.append(row)
    for c, v in zip(t.rows[-1].cells, values):
        for p in c.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        runs = c.paragraphs[0].runs
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        if runs:
            runs[0].text = v
        else:
            c.paragraphs[0].add_run(v)


for t in d.tables:
    head = [c.text for c in t.rows[0].cells]
    if any("Anupria" in c.text for r in t.rows for c in r.cells):
        continue
    if head == ["Sr. No.", "Student Name", "PRN"]:
        add_row(t, ["3", *NEW])
    elif head == ["Student Name", "PRN"]:
        add_row(t, list(NEW))
    elif head == ["Student Name", "PRN", "Signature"]:
        add_row(t, [*NEW, ""])
    elif head[0] == "Student Name and PRN":
        t.rows[2].cells[1].paragraphs[0].runs[0].text = "Network emulator and experiments; presentation"
        t.rows[2].cells[2].paragraphs[0].runs[0].text = "mpft/emulator.py, experiments/, results/"
        add_row(t, [f"{NEW[0]}, PRN {NEW[1]}", "Wireshark dissector and capture steps; web interface views and screenshots; test report",
                    "wireshark/, docs/screenshots/, docs/TEST_REPORT.md"])
d.save(R)

P = "docs/submission/MPFT_Presentation.pptx"
pr = pptx.Presentation(P)
s0 = pr.slides[0]
box = next(sh for sh in s0.shapes if sh.name == "Rectangle 6")
if "Anupria" not in box.text_frame.text:
    p0 = box.text_frame.paragraphs[1]._p
    p0.addnext(copy.deepcopy(p0))
    last = box.text_frame.paragraphs[2]
    last.runs[0].text, last.runs[1].text = NEW[0], f" – PRN {NEW[1]}"
    box.height = Emu(int(box.height * 1.45))
s7 = pr.slides[7]
a = next(sh for sh in s7.shapes if sh.name == "TextBox 21")
b = next(sh for sh in s7.shapes if sh.name == "TextBox 22")
if "Anupria" not in "".join(sh.text_frame.text for sh in s7.shapes if sh.has_text_frame):
    c = copy.deepcopy(b._element)
    b._element.addnext(c)
    c = next(sh for sh in s7.shapes if sh._element is c)
    w, gap, x0 = Emu(2600000), Emu(120000), a.left
    for i, sh in enumerate((a, b, c)):
        sh.left, sh.width = Emu(x0 + i * (w + gap)), w
    a.text_frame.paragraphs[1].runs[0].text = "Protocol, sender, receiver, resume, web UI"
    b.text_frame.paragraphs[1].runs[0].text = "Emulator, experiments, presentation"
    c.text_frame.paragraphs[0].runs[0].text = NEW[0]
    c.text_frame.paragraphs[1].runs[0].text = "Wireshark, dashboard views, test report"
pr.save(P)
print("ok")
